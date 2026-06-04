import pandas as pd
import openpyxl
import os
import streamlit as st
import plotly.express as px
import re
import folium
import numpy as np
from dotenv import load_dotenv

from scipy.spatial import ConvexHull
from folium.plugins import HeatMap, MarkerCluster, MiniMap, Fullscreen
from streamlit_folium import st_folium

load_dotenv()


# =========================================================
# 配置与常量管理
# =========================================================
class Config:
    PAGE_TITLE = "大型充放电站选址定容规划"
    PAGE_ICON = "🗺️"
    LAYOUT = "wide"

    MAP_CENTER = [22.93, 113.45]
    MAP_ZOOM = 11
    MAP_HEIGHT = 760

    API_KEY = os.getenv("THUNDERFOREST_API_KEY")
    TILE_URL = "https://tile.thunderforest.com/atlas/{{z}}/{{x}}/{{y}}.png?apikey={api_key}"

    # TILE_URL = "https://server.arcgisonline.com/ArcGIS/rest/services/World_Imagery/MapServer/tile/{z}/{y}/{x}"
    # TILE_URL = "https://{s}.basemaps.cartocdn.com/dark_all/{z}/{x}/{y}{r}.png"
    # TILE_URL = "https://tile.thunderforest.com/transport-dark/{{z}}/{{x}}/{{y}}.png?apikey={api_key}"
    # TILE_URL = "https://tile.thunderforest.com/transport/{{z}}/{{x}}/{{y}}.png?apikey={api_key}"

    FIXED_EXCEL = "大型案例.xlsx"

    COLORS = {
        "phase1": "#00cfff",
        "phase2": "#ff4d4f",
        "expansion": "#ffd666",
        "selected": "#FFD700",
        "station": "#00e5ff",

        "station_level4": "#90EE90",  # 四级：3-50个充电桩 - 浅绿色
        "station_level3": "#FFA500",  # 三级：51-150个充电桩 - 橙色
        "station_level2": "#FF4500",  # 二级：151-300个充电桩 - 橙红色
        "station_level1": "#DC143C"  # 一级：300个以上充电桩 - 深红色
    }

    HEAT_GRADIENTS = {
        "phase1": {
            0.2: '#00ffff',
            0.4: '#00bfff',
            0.6: '#1e90ff',
            0.8: '#4169e1',
            1.0: '#7b68ee'
        },
        "phase2": {
            0.2: '#ffcccc',
            0.4: '#ff9999',
            0.6: '#ff6666',
            0.8: '#ff3333',
            1.0: '#cc0000'
        },
        "expansion": {
            0.2: '#fff7bc',
            0.4: '#fee391',
            0.6: '#fec44f',
            0.8: '#fe9929',
            1.0: '#d95f0e'
        }
    }

    POLYGON_COLORS = [
        '#00FFFF',
        '#00BFFF',
        '#1E90FF',
        '#4169E1',
        '#7B68EE',
        '#9370DB',
        '#BA55D3',
        '#20B2AA'
    ]

    EXCEL_COLUMNS = {
        "station_id": 0,  # A列：充电站ID
        "pdn_node": 5,  # F列：PDN节点
        "capacity_s1": 6,  # G列：CapacityS1（阶段一充电桩数）
        "capacity_s2": 7,  # H列：CapacityS2（阶段二充电桩数）
        "selected": 8,  # I列：选中标记
    }

    # 单桩额定功率（kW）
    CHARGER_POWER = 7.7

    STATION_LEVELS = {
        "level4": {"min": 3, "max": 50, "color": "station_level4", "name": "四级充电站"},
        "level3": {"min": 51, "max": 150, "color": "station_level3", "name": "三级充电站"},
        "level2": {"min": 151, "max": 300, "color": "station_level2", "name": "二级充电站"},
        "level1": {"min": 301, "max": float('inf'), "color": "station_level1", "name": "一级充电站"}
    }


# =========================================================
# 数据加载与处理
# =========================================================
class DataLoader:

    @staticmethod
    @st.cache_data
    def load_data():

        script_dir = os.path.dirname(os.path.abspath(__file__))

        fixed_path = os.path.join(
            script_dir,
            Config.FIXED_EXCEL
        )

        excel_path = fixed_path

        wb = openpyxl.load_workbook(
            excel_path,
            data_only=True
        )

        # =====================================================
        # 读取服务区与充电站总表
        # =====================================================
        ws_main = wb["充电服务区分区情况"]

        charging_stations_base = []

        for row in ws_main.iter_rows(
                min_row=2,
                values_only=True
        ):

            try:
                station_id, lng, lat, cluster = row[:4]

                if all(v is not None for v in [station_id, lng, lat, cluster]):
                    charging_stations_base.append({
                        "充电站ID": int(station_id),
                        "lon": float(lng),
                        "lat": float(lat),
                        "所属服务区": int(cluster)
                    })

            except:
                continue

        stations_base_df = pd.DataFrame(charging_stations_base)

        # =====================================================
        # 读取每个服务区的充电站详细信息
        # =====================================================
        all_stations_detail = []
        service_info_list = []

        service_sheets = []

        for sheet_name in wb.sheetnames:

            if re.match(r"服务区\d+$", sheet_name):
                service_sheets.append(sheet_name)

        for sheet_name in service_sheets:

            try:

                ws = wb[sheet_name]

                rows = list(ws.iter_rows(values_only=True))

                if len(rows) < 2:
                    continue

                # =================================================
                # 服务区基础信息
                # =================================================
                basic_row = rows[1]

                area_id = int(basic_row[0])
                area_lng = float(basic_row[1])
                area_lat = float(basic_row[2])
                pdn_node = str(basic_row[Config.EXCEL_COLUMNS["pdn_node"]]).strip() if basic_row[Config.EXCEL_COLUMNS[
                    "pdn_node"]] is not None else "未配置"

                # =================================================
                # 读取每个充电站的详细信息
                # =================================================
                area_stations = []
                area_total_s1 = 0
                area_total_s2 = 0

                for r in rows[4:]:

                    try:

                        # 提取充电站ID
                        if r[Config.EXCEL_COLUMNS["station_id"]] is None:
                            continue
                        station_id = int(r[Config.EXCEL_COLUMNS["station_id"]])

                        # 读取选中标记
                        selected_flag = (
                            str(r[Config.EXCEL_COLUMNS["selected"]]).strip()
                            if r[Config.EXCEL_COLUMNS["selected"]] is not None else ""
                        )

                        if selected_flag != "选中":
                            continue

                        capacity_s1 = (
                            int(r[Config.EXCEL_COLUMNS["capacity_s1"]])
                            if r[Config.EXCEL_COLUMNS["capacity_s1"]] is not None else 0
                        )

                        capacity_s2 = (
                            int(r[Config.EXCEL_COLUMNS["capacity_s2"]])
                            if r[Config.EXCEL_COLUMNS["capacity_s2"]] is not None else 0
                        )

                        all_stations_detail.append({
                            "充电站ID": station_id,
                            "CapacityS1": capacity_s1,
                            "CapacityS2": capacity_s2,
                            "CapacityExpansion": capacity_s2 - capacity_s1
                        })

                        # 累加服务区总容量
                        area_total_s1 += capacity_s1
                        area_total_s2 += capacity_s2
                        area_stations.append(station_id)

                    except Exception as e:
                        print(f"读取行出错: {e}")
                        continue

                # =================================================
                # 保存服务区汇总信息
                # =================================================
                service_info_list.append({
                    "服务区ID": area_id,
                    "lon": area_lng,
                    "lat": area_lat,
                    "充电站数量": len(area_stations),
                    "充电桩数_阶段1": area_total_s1,
                    "充电桩数_阶段2": area_total_s2,
                    "充电桩扩容数": area_total_s2 - area_total_s1,
                    "PDN节点": pdn_node
                })

            except Exception as e:
                print(f"读取 {sheet_name} 出错: {e}")
                continue

        # =====================================================
        # 合并充电站基础信息与详细信息
        # =====================================================
        stations_detail_df = pd.DataFrame(all_stations_detail)

        selected_charging_stations = pd.merge(
            stations_base_df,
            stations_detail_df,
            on="充电站ID",
            how="inner"
        ).reset_index(drop=True)

        all_service_areas = pd.DataFrame(
            service_info_list
        ).sort_values(
            "服务区ID"
        ).reset_index(drop=True)

        return {
            "all_charging_stations": stations_base_df,
            "selected_charging_stations": selected_charging_stations,
            "all_service_areas": all_service_areas
        }


# =========================================================
# 地图生成与交互
# =========================================================
class MapGenerator:

    @staticmethod
    def create_map(display_df, filtered_stations_df, view_mode, center=None):

        display_config = MapGenerator._get_display_config(view_mode)

        map_center = center if center is not None else Config.MAP_CENTER

        m = folium.Map(
            location=map_center,
            zoom_start=Config.MAP_ZOOM,
            tiles=None,
            control_scale=True,
            prefer_canvas=True
        )

        # 添加热力图层级
        folium.map.CustomPane(
            "heatmapPane",
            z_index=200,
            pointer_events=False
        ).add_to(m)

        # 添加底图
        MapGenerator._add_tile_layer(m)

        # 添加充电站位置点
        MapGenerator._add_station_layer(
            m,
            filtered_stations_df,
            view_mode
        )

        # 添加服务区边界
        MapGenerator._add_polygon_layer(
            m,
            filtered_stations_df
        )

        # 添加热力图
        MapGenerator._add_heatmap(
            m,
            filtered_stations_df,
            display_config["heat_gradient"]
        )

        # 添加服务区中心聚合层
        MapGenerator._add_service_area_markers(
            m,
            display_df,
            display_config["color_main"]
        )

        # 添加充电站等级图例
        MapGenerator._add_station_level_legend(m)

        # 添加小地图和全屏控件
        MapGenerator._add_map_controls(m)

        # 添加图层控制
        folium.LayerControl().add_to(m)

        return m

    @staticmethod
    def _get_display_config(view_mode):

        config_map = {
            "优化阶段一": {
                "color_main": Config.COLORS["phase1"],
                "heat_gradient": Config.HEAT_GRADIENTS["phase1"]
            },
            "优化阶段二": {
                "color_main": Config.COLORS["phase2"],
                "heat_gradient": Config.HEAT_GRADIENTS["phase2"]
            },
            "扩容演化": {
                "color_main": Config.COLORS["expansion"],
                "heat_gradient": Config.HEAT_GRADIENTS["expansion"]
            }
        }

        return config_map.get(
            view_mode,
            config_map["优化阶段一"]
        )

    @staticmethod
    def _add_tile_layer(m):

        tile_url = Config.TILE_URL.format(
            api_key=Config.API_KEY
        )

        folium.TileLayer(
            tiles=tile_url,
            attr="Thunderforest",
            name="Atlas",
            overlay=False,
            control=True
        ).add_to(m)

    @staticmethod
    def _add_station_layer(m, filtered_stations_df, view_mode):
        station_layer = folium.FeatureGroup(
            name="充电站位置",
            show=True
        )

        selected_area = st.session_state.selected_area

        if view_mode == "优化阶段一":
            capacity_field = "CapacityS1"
        elif view_mode == "优化阶段二":
            capacity_field = "CapacityS2"
        else:
            capacity_field = "CapacityExpansion"

        for _, row in filtered_stations_df.iterrows():

            is_selected = (
                    selected_area is not None
                    and row["所属服务区"] == selected_area
            )

            # 获取该充电站当前阶段的充电桩数量
            charger_count = row[capacity_field]

            if is_selected:
                # 选中状态保持金色高亮
                radius, opacity, color = (
                    3.5,
                    0.92,
                    Config.COLORS["selected"]
                )
            else:
                if charger_count >= Config.STATION_LEVELS["level1"]["min"]:
                    color = Config.COLORS[Config.STATION_LEVELS["level1"]["color"]]
                    radius, opacity = 3.2, 0.7
                elif charger_count >= Config.STATION_LEVELS["level2"]["min"]:
                    color = Config.COLORS[Config.STATION_LEVELS["level2"]["color"]]
                    radius, opacity = 2.9, 0.6
                elif charger_count >= Config.STATION_LEVELS["level3"]["min"]:
                    color = Config.COLORS[Config.STATION_LEVELS["level3"]["color"]]
                    radius, opacity = 2.6, 0.5
                elif charger_count >= Config.STATION_LEVELS["level4"]["min"]:
                    color = Config.COLORS[Config.STATION_LEVELS["level4"]["color"]]
                    radius, opacity = 2.3, 0.4
                else:
                    color = Config.COLORS["station"]
                    radius, opacity = 2.0, 0.3

            folium.CircleMarker(
                location=[row["lat"], row["lon"]],
                radius=radius,
                stroke=False,
                fill=True,
                fill_color=color,
                fill_opacity=opacity,
                pane="markerPane",
                interactive=False
            ).add_to(station_layer)

        station_layer.add_to(m)

    @staticmethod
    def _add_polygon_layer(m, filtered_stations_df):

        polygon_layer = folium.FeatureGroup(
            name="服务区边界",
            show=False
        )

        station_groups = filtered_stations_df.groupby(
            "所属服务区"
        )

        selected_area = st.session_state.selected_area

        for idx, (area_id, group) in enumerate(station_groups):

            if len(group) < 3:
                continue

            points = group[["lon", "lat"]].values

            try:

                hull = ConvexHull(points)

                hull_points = points[hull.vertices]

                polygon_coords = [
                    [p[1], p[0]]
                    for p in hull_points
                ]

                if area_id == selected_area:
                    border_color = Config.COLORS["selected"]
                    border_weight = 5
                    border_opacity = 1.0
                else:
                    border_color = Config.POLYGON_COLORS[
                        idx % len(Config.POLYGON_COLORS)
                        ]
                    border_weight = 2.5
                    border_opacity = 0.7

                folium.Polygon(
                    locations=polygon_coords,
                    color=border_color,
                    weight=border_weight,
                    opacity=border_opacity,
                    fill=False,
                    tooltip=f"服务区 {area_id}<br>充电站数量: {len(group)}"
                ).add_to(polygon_layer)

            except:
                pass

        polygon_layer.add_to(m)

    @staticmethod
    def _add_heatmap(m, filtered_stations_df, heat_gradient):

        heat_data = [
            [row["lat"], row["lon"], 1]
            for _, row in filtered_stations_df.iterrows()
        ]

        heatmap = HeatMap(
            heat_data,
            radius=16,
            blur=10,
            min_opacity=0.12,
            gradient=heat_gradient,
            overlay=True,
            control=True,
            show=True,
            name="热力图"
        )

        heatmap.add_to(m)

        heatmap.options["pane"] = "heatmapPane"

    @staticmethod
    def _add_service_area_markers(
            m,
            display_df,
            color_main
    ):

        marker_cluster = MarkerCluster(
            name="服务区中心",
            spiderfyOnMaxZoom=True,
            showCoverageOnHover=False,
            zoomToBoundsOnClick=True,
            disableClusteringAtZoom=1
        ).add_to(m)

        selected_area = st.session_state.selected_area

        for _, row in display_df.iterrows():

            area_id = int(row["服务区ID"])

            jitter_seed = area_id * 0.00001

            jlat = (
                    row["lat"]
                    + np.sin(jitter_seed * 1000) * 0.00015
            )

            jlng = (
                    row["lon"]
                    + np.cos(jitter_seed * 1000) * 0.00015
            )

            radius = max(
                6,
                row["显示充电桩数"] / 120
            )

            if selected_area == area_id:
                border_color = Config.COLORS["selected"]
                border_width = 5
            else:
                border_color = "white"
                border_width = 2

            tooltip_html = MapGenerator._create_tooltip_html(
                row
            )

            folium.CircleMarker(
                location=[jlat, jlng],
                radius=radius,
                color=border_color,
                weight=border_width,
                fill=True,
                fill_color=color_main,
                fill_opacity=0.95,
                tooltip=folium.Tooltip(
                    tooltip_html,
                    sticky=True
                )
            ).add_to(marker_cluster)

    @staticmethod
    @st.cache_data
    def _get_area_positions(display_df):
        area_positions = {}

        for _, row in display_df.iterrows():
            area_id = int(row["服务区ID"])

            jitter_seed = area_id * 0.00001

            jlat = (
                    row["lat"]
                    + np.sin(jitter_seed * 1000) * 0.00015
            )

            jlng = (
                    row["lon"]
                    + np.cos(jitter_seed * 1000) * 0.00015
            )

            area_positions[area_id] = (jlat, jlng)

        return area_positions

    @staticmethod
    def _create_tooltip_html(row):
        return f"""
        <div style="width:280px; font-family:Arial; color:#2c3e50;">
            <h4 style="margin-bottom:10px; color:#1677ff;">
                📍 服务区 {row['服务区ID']}
            </h4>
            <b>PDN节点：</b>{row['PDN节点']}<br>
            <b>当前阶段充电桩总数：</b>
            {row['显示充电桩数']} 个<br>

            <b>包含充电站数：</b>
            {row['充电站数量']} 个<br>

            <b>总充电功率：</b>
            {row['总充电功率']:.1f} kW<br>
        </div>
        """

    @staticmethod
    def _add_station_level_legend(m):
        legend_html = '''
        <div style="
            position: fixed;
            top: 20px;
            left: 20px;
            z-index: 1000;
            background-color: rgba(0, 0, 0, 0.7);
            color: white;
            padding: 12px 15px;
            border-radius: 8px;
            font-family: Arial;
            font-size: 13px;
            box-shadow: 0 2px 10px rgba(0,0,0,0.3);
        ">
            <div style="font-weight: bold; margin-bottom: 8px; text-align: center;">
                充电站等级
            </div>
            <div style="display: flex; align-items: center; margin-bottom: 5px;">
                <div style="
                    width: 14px;
                    height: 14px;
                    border-radius: 50%;
                    background-color: #DC143C;
                    margin-right: 8px;
                "></div>
                <span>一级充电站 (≥301桩)</span>
            </div>
            <div style="display: flex; align-items: center; margin-bottom: 5px;">
                <div style="
                    width: 14px;
                    height: 14px;
                    border-radius: 50%;
                    background-color: #FF4500;
                    margin-right: 8px;
                "></div>
                <span>二级充电站 (151-300桩)</span>
            </div>
            <div style="display: flex; align-items: center; margin-bottom: 5px;">
                <div style="
                    width: 14px;
                    height: 14px;
                    border-radius: 50%;
                    background-color: #FFA500;
                    margin-right: 8px;
                "></div>
                <span>三级充电站 (51-150桩)</span>
            </div>
            <div style="display: flex; align-items: center;">
                <div style="
                    width: 14px;
                    height: 14px;
                    border-radius: 50%;
                    background-color: #90EE90;
                    margin-right: 8px;
                "></div>
                <span>四级充电站 (3-50桩)</span>
            </div>
        </div>
        '''

        m.get_root().html.add_child(folium.Element(legend_html))

    @staticmethod
    def _add_map_controls(m):

        minimap_tiles = folium.TileLayer(
            tiles=Config.TILE_URL.format(
                api_key=Config.API_KEY
            ),
            attr="Thunderforest"
        )

        MiniMap(
            tile_layer=minimap_tiles,
            toggle_display=True,
            width=150,
            height=150,
            zoom_level_offset=-5
        ).add_to(m)

        Fullscreen(position="topright").add_to(m)

    @staticmethod
    def handle_map_click(current_click, display_df):
        if current_click is None:
            return False

        try:
            click_lat = float(current_click["lat"])
            click_lng = float(current_click["lng"])

            area_positions = MapGenerator._get_area_positions(display_df)

            min_dist = float('inf')
            nearest_area = None

            for area_id, pos in area_positions.items():
                s_lat, s_lng = pos
                dist = (s_lat - click_lat) ** 2 + (s_lng - click_lng) ** 2
                if dist < min_dist:
                    min_dist = dist
                    nearest_area = area_id

            if min_dist < 0.000005 and st.session_state.selected_area != nearest_area:
                st.session_state.selected_area = nearest_area
                return True

            return False

        except:
            return False


# =========================================================
# UI组件与图表
# =========================================================
class UIComponents:

    @staticmethod
    def setup_page():
        st.set_page_config(
            page_title=Config.PAGE_TITLE,
            page_icon=Config.PAGE_ICON,
            layout=Config.LAYOUT
        )

        st.markdown("""
        <style>
        .stApp {
            background-color: #f5f7fa;
        }
        .dashboard-wrapper {
            background: #ffffff;
            border-radius: 16px;
            padding: 20px;
            margin-bottom: 12px;
            box-shadow: 0 2px 12px rgba(0, 0, 0, 0.06);
            border: 1px solid #ebeef5;
        }
        /* 标题栏 */
        .dashboard-module-title {
            font-size: 18px;
            font-weight: 700;
            color: #1d2129;
            margin: 32px 0 18px 0;
            padding-left: 12px;
            border-left: 4px solid #1677ff;
            line-height: 1.4;
        }

        .main-kpi-card {
            background: linear-gradient(145deg, #ffffff, #f0f7ff);
            border-radius: 14px;
            padding: 22px 18px;
            height: 100%;
            box-shadow: 0 4px 16px rgba(22, 119, 255, 0.12), 
                        0 1px 4px rgba(0, 0, 0, 0.05);
            border: 1px solid rgba(22, 119, 255, 0.15);
            transition: all 0.25s ease;
            position: relative;
            overflow: hidden;
        }
        .main-kpi-card:hover {
            box-shadow: 0 6px 22px rgba(22, 119, 255, 0.18), 
                        0 2px 8px rgba(0, 0, 0, 0.08);
            transform: translateY(-2px);
        }
        .main-card-icon {
            font-size: 30px;
            color: #1677ff;
            margin-bottom: 10px;
        }
        .main-card-label {
            font-size: 14px;
            color: #6e7681;
            font-weight: 500;
            margin-bottom: 6px;
        }
        .main-card-value {
            font-size: 36px;
            font-weight: 800;
            color: #1d2129;
            line-height: 1.1;
        }
        .main-card-unit {
            font-size: 16px;
            color: #1677ff;
            font-weight: 600;
            margin-left: 6px;
        }
        .main-card-desc {
            font-size: 12px;
            color: #86909c;
            margin-top: 6px;
        }

        .sub-kpi-card {
            background: #ffffff;
            border-radius: 12px;
            padding: 18px 14px;
            height: 100%;
            box-shadow: 0 2px 8px rgba(0, 0, 0, 0.06);
            border: 1px solid #ebeef5;
            transition: all 0.25s ease;
        }
        .sub-kpi-card:hover {
            box-shadow: 0 4px 12px rgba(0, 0, 0, 0.1);
        }
        .sub-card-icon {
            font-size: 22px;
            color: #86909c;
            margin-bottom: 6px;
        }
        .sub-card-label {
            font-size: 13px;
            color: #6e7681;
            font-weight: 500;
            margin-bottom: 4px;
        }
        .sub-card-value {
            font-size: 26px;
            font-weight: 700;
            color: #1d2129;
        }
        .level-kpi-card {
            background: #ffffff;
            border-radius: 12px;
            padding: 16px 14px;
            height: 100%;
            box-shadow: 0 2px 8px rgba(0, 0, 0, 0.06);
            border-left: 5px solid;
            border-top: 1px solid #ebeef5;
            border-right: 1px solid #ebeef5;
            border-bottom: 1px solid #ebeef5;
            transition: all 0.25s ease;
        }
        .level-kpi-card:hover {
            box-shadow: 0 4px 14px rgba(0, 0, 0, 0.1);
        }
        .level-card-name {
            font-size: 14px;
            color: #1d2129;
            font-weight: 600;
            margin-bottom: 4px;
        }
        .level-card-count {
            font-size: 28px;
            font-weight: 800;
            margin: 4px 0;
        }
        .level-card-range {
            font-size: 12px;
            color: #86909c;
            margin-bottom: 8px;
        }
        .stProgress {
            margin-top: 0 !important;
            margin-bottom: 0 !important;
        }
        .stProgress > div > div > div > div {
            border-radius: 10px !important;
            height: 8px !important;
        }
        </style>
        """, unsafe_allow_html=True)

        st.title(f"{Config.PAGE_ICON} {Config.PAGE_TITLE}")
        st.caption("充电站选址与容量优化可视化系统")

    @staticmethod
    def init_session_state():
        if "selected_area" not in st.session_state:
            st.session_state.selected_area = None

        if "last_processed_click" not in st.session_state:
            st.session_state.last_processed_click = None

        if "display_mode" not in st.session_state:
            st.session_state.display_mode = "chart"

    @staticmethod
    def show_kpi_metrics(
            filtered_stations_df,
            display_df,
            view_mode
    ):
        total_stations = len(filtered_stations_df)
        total_power = display_df['总充电功率'].sum()
        total_areas = len(display_df)
        total_chargers = display_df['显示充电桩数'].sum()

        st.markdown('<div class="dashboard-module-title">📊 系统核心指标总览</div>', unsafe_allow_html=True)

        col_main1, col_main2, col_sub1, col_sub2 = st.columns([2, 2, 1, 1])

        with col_main1:
            st.markdown(f"""
            <div class="main-kpi-card">
                <div class="main-card-icon">🏭</div>
                <div class="main-card-label">当前阶段充电站总数</div>
                <div class="main-card-value">{total_stations}<span class="main-card-unit">座</span></div>
                <div class="main-card-desc">优化布局内有效充电站数量</div>
            </div>
            """, unsafe_allow_html=True)

        with col_main2:
            st.markdown(f"""
            <div class="main-kpi-card">
                <div class="main-card-icon">⚡</div>
                <div class="main-card-label">系统总充电功率</div>
                <div class="main-card-value">{total_power:.1f}<span class="main-card-unit">kW</span></div>
                <div class="main-card-desc">全部充电桩额定总输出功率</div>
            </div>
            """, unsafe_allow_html=True)

        with col_sub1:
            st.markdown(f"""
            <div class="sub-kpi-card">
                <div class="sub-card-icon">📍</div>
                <div class="sub-card-label">服务区数量</div>
                <div class="sub-card-value">{total_areas}</div>
            </div>
            """, unsafe_allow_html=True)

        with col_sub2:
            st.markdown(f"""
            <div class="sub-kpi-card">
                <div class="sub-card-icon">🔌</div>
                <div class="sub-card-label">充电桩总数</div>
                <div class="sub-card-value">{total_chargers:.0f}</div>
            </div>
            """, unsafe_allow_html=True)

        st.markdown('<div class="dashboard-module-title">🏆 充电站等级结构分布</div>', unsafe_allow_html=True)

        if view_mode == "优化阶段一":
            capacity_field = "CapacityS1"
        elif view_mode == "优化阶段二":
            capacity_field = "CapacityS2"
        else:
            capacity_field = "CapacityExpansion"
        level1_count = len(
            filtered_stations_df[filtered_stations_df[capacity_field] >= Config.STATION_LEVELS["level1"]["min"]])
        level2_count = len(
            filtered_stations_df[(filtered_stations_df[capacity_field] >= Config.STATION_LEVELS["level2"]["min"]) &
                                 (filtered_stations_df[capacity_field] < Config.STATION_LEVELS["level1"]["min"])])
        level3_count = len(
            filtered_stations_df[(filtered_stations_df[capacity_field] >= Config.STATION_LEVELS["level3"]["min"]) &
                                 (filtered_stations_df[capacity_field] < Config.STATION_LEVELS["level2"]["min"])])
        level4_count = len(
            filtered_stations_df[(filtered_stations_df[capacity_field] >= Config.STATION_LEVELS["level4"]["min"]) &
                                 (filtered_stations_df[capacity_field] < Config.STATION_LEVELS["level3"]["min"])])

        level_data = [
            {
                "name": "一级充电站",
                "count": level1_count,
                "color": Config.COLORS["station_level1"],
                "range": "≥ 301 个充电桩"
            },
            {
                "name": "二级充电站",
                "count": level2_count,
                "color": Config.COLORS["station_level2"],
                "range": "151 ~ 300 个充电桩"
            },
            {
                "name": "三级充电站",
                "count": level3_count,
                "color": Config.COLORS["station_level3"],
                "range": "51 ~ 150 个充电桩"
            },
            {
                "name": "四级充电站",
                "count": level4_count,
                "color": Config.COLORS["station_level4"],
                "range": "3 ~ 50 个充电桩"
            }
        ]

        col_l1, col_l2, col_l3, col_l4 = st.columns(4)
        cols = [col_l1, col_l2, col_l3, col_l4]

        for i, level in enumerate(level_data):
            with cols[i]:

                st.markdown(f"""
                <div class="level-kpi-card" style="border-left-color: {level['color']};">
                    <div class="level-card-name">{level['name']}</div>
                    <div class="level-card-count" style="color:{level['color']};">{level['count']}</div>
                    <div class="level-card-range">{level['range']}</div>
                </div>
                """, unsafe_allow_html=True)

    @staticmethod
    def create_station_level_chart(filtered_stations_df, view_mode):

        # =========================================================
        # 阶段配置
        # =========================================================
        if view_mode == "优化阶段一":
            capacity_field = "CapacityS1"
            stage_name = "阶段一"
        elif view_mode == "优化阶段二":
            capacity_field = "CapacityS2"
            stage_name = "阶段二"
        else:
            capacity_field = "CapacityExpansion"
            stage_name = "扩容规划"

        # =========================================================
        # 等级统计
        # =========================================================
        level_stats = []

        level_definitions = [
            ("level1", "一级充电站", "#00B4D8"),
            ("level2", "二级充电站", "#38B000"),
            ("level3", "三级充电站", "#FF9F1C"),
            ("level4", "四级充电站", "#E63946")
        ]

        for level_key, level_name, color in level_definitions:
            config = Config.STATION_LEVELS[level_key]

            if level_key == "level1":
                count = len(
                    filtered_stations_df[
                        filtered_stations_df[capacity_field] >= config["min"]
                        ]
                )
                range_text = f"≥{config['min']}桩"
            else:
                count = len(
                    filtered_stations_df[
                        (filtered_stations_df[capacity_field] >= config["min"]) &
                        (filtered_stations_df[capacity_field] < config["max"])
                        ]
                )
                range_text = f"{config['min']}-{config['max']}桩"

            level_stats.append({
                "等级": level_name,
                "容量范围": range_text,
                "数量": count,
                "颜色": color
            })

        stats_df = pd.DataFrame(level_stats)
        stats_df = stats_df[stats_df["数量"] > 0]

        total_stations = stats_df["数量"].sum()

        MIN_PERCENT = 0.02
        MAX_PULL_DISTANCE = 0.25
        BASE_PULL_DISTANCE = 0.05

        stats_df["原始数量"] = stats_df["数量"]
        stats_df["原始占比"] = stats_df["数量"] / total_stations

        stats_df["绘制数值"] = stats_df.apply(
            lambda row: max(row["数量"], total_stations * MIN_PERCENT),
            axis=1
        )

        stats_df["pull_distance"] = stats_df["原始占比"].apply(
            lambda p: BASE_PULL_DISTANCE + (MAX_PULL_DISTANCE - BASE_PULL_DISTANCE) * (1 - p / MIN_PERCENT)
            if p < MIN_PERCENT else BASE_PULL_DISTANCE
        )

        stats_df["标签文本"] = stats_df.apply(
            lambda row: f"{row['等级']}<br>{row['原始占比']:.1%}",
            axis=1
        )

        fig = px.pie(
            stats_df,
            values="绘制数值",
            names="等级",
            color="等级",
            color_discrete_map=dict(zip(stats_df["等级"], stats_df["颜色"])),
            hover_data=None
        )

        fig.update_traces(
            text=stats_df["标签文本"],
            textposition='outside',
            textinfo='text',
            textfont=dict(
                size=12,
                color="#2D3748",
                family="Microsoft YaHei",
                weight="bold"
            ),
            pull=stats_df["pull_distance"].tolist(),
            rotation=90,
            sort=False,
            showlegend=True,
            hoverinfo="none",
            hovertemplate=None,
            customdata=None,
            marker=dict(
                colors=stats_df["颜色"],
                line=dict(
                    color="#FFFFFF",
                    width=2.5
                )
            )
        )

        fig.update_layout(
            height=550,
            margin=dict(l=50, r=50, t=80, b=100),
            title=dict(
                text=(
                    f"<span style='font-size:13px;color:#718096;'>"
                    f"{stage_name} · 总计 {total_stations:,} 座"
                    f"</span>"
                ),
                x=0.5,
                y=0.96,
                xanchor="center",
                yanchor="top",
                font=dict(size=24, color="#1A202C")
            ),
            legend=dict(
                orientation="h",
                yanchor="bottom",
                y=-0.25,
                xanchor="center",
                x=0.5,
                bgcolor="#F7FAFC",
                bordercolor="#E2E8F0",
                borderwidth=1,
                font=dict(size=13, color="#2D3748"),
                itemwidth=30,
                itemsizing="constant",
                tracegroupgap=0,
                entrywidth=0,
                entrywidthmode="pixels"
            ),
            transition=dict(duration=800, easing="cubic-in-out"),
            paper_bgcolor="#FFFFFF",
            plot_bgcolor="#FFFFFF"
        )

        st.plotly_chart(
            fig,
            use_container_width=True,
            config={"displayModeBar": False, "responsive": True}
        )

    @staticmethod
    def show_selected_area_details(
            all_service_areas,
            filtered_stations_df
    ):
        st.markdown("### 🎯 当前选中服务区")

        if st.session_state.selected_area is None:
            st.info("点击地图中的服务区以查看该服务区内部充电站列表。")
            return

        selected_row = all_service_areas[
            all_service_areas["服务区ID"] == st.session_state.selected_area
            ]

        if len(selected_row) == 0:
            st.warning("未找到该服务区信息。")
            return

        row = selected_row.iloc[0]
        pdn_node = row["PDN节点"]

        st.success(f"当前选中：服务区 {row['服务区ID']}")

        stations_in_area = filtered_stations_df[
            filtered_stations_df["所属服务区"] == st.session_state.selected_area
            ].copy()

        power_s1 = row['充电桩数_阶段1'] * Config.CHARGER_POWER
        power_s2 = row['充电桩数_阶段2'] * Config.CHARGER_POWER
        power_expansion = row['充电桩扩容数'] * Config.CHARGER_POWER

        st.markdown(
            f"""
            <div style="
                padding: 1rem;
                border-radius: 0.5rem;
                background-color: #e6f7ff;
                border-left: 5px solid #1890ff;
                margin-bottom: 1rem;
            ">
                <div style="color: #0050b3; font-weight: 500;">
                    📶 关联PDN节点：{pdn_node}<br>
                    ℹ️ 当前阶段包含 {len(stations_in_area)} 个充电站<br>
                    🔵 阶段一总充电桩数：{row['充电桩数_阶段1']} 个（{power_s1:.1f} kW）<br>
                    🔴 阶段二总充电桩数：{row['充电桩数_阶段2']} 个（{power_s2:.1f} kW）<br>
                    🟡 总新增充电桩数：{row['充电桩扩容数']} 个（{power_expansion:.1f} kW）
                </div>
            </div>
            """,
            unsafe_allow_html=True
        )

        stations_display_df = stations_in_area[
            ["充电站ID", "lon", "lat", "CapacityS1", "CapacityS2", "CapacityExpansion"]
        ].copy()

        stations_display_df.columns = [
            "充电站ID", "经度", "纬度", "阶段一充电桩数", "阶段二充电桩数", "新增充电桩数"
        ]

        st.dataframe(
            stations_display_df,
            use_container_width=True,
            height=400,
            hide_index=True
        )


def main():
    UIComponents.setup_page()

    UIComponents.init_session_state()

    data = DataLoader.load_data()

    selected_charging_stations = data["selected_charging_stations"]
    all_service_areas = data["all_service_areas"]

    view_mode = st.radio(
        "优化阶段",
        [
            "优化阶段一",
            "优化阶段二",
            "扩容演化"
        ],
        horizontal=True
    )

    if view_mode == "优化阶段一":
        phase_stations = selected_charging_stations[
            selected_charging_stations["CapacityS1"] > 0
            ].copy()

        display_df = all_service_areas[
            all_service_areas["充电桩数_阶段1"] > 0
            ].copy()

        display_df["显示充电桩数"] = display_df["充电桩数_阶段1"]
        display_df["总充电功率"] = display_df["显示充电桩数"] * Config.CHARGER_POWER
        chart_title = "优化阶段一布局"

    elif view_mode == "优化阶段二":
        phase_stations = selected_charging_stations[
            selected_charging_stations["CapacityS2"] > 0
            ].copy()

        display_df = all_service_areas[
            all_service_areas["充电桩数_阶段2"] > 0
            ].copy()

        display_df["显示充电桩数"] = display_df["充电桩数_阶段2"]
        display_df["总充电功率"] = display_df["显示充电桩数"] * Config.CHARGER_POWER
        chart_title = "优化阶段二布局"

    else:
        phase_stations = selected_charging_stations[
            selected_charging_stations["CapacityExpansion"] > 0
            ].copy()

        display_df = all_service_areas[
            all_service_areas["充电桩扩容数"] > 0
            ].copy()

        display_df["显示充电桩数"] = display_df["充电桩扩容数"]
        display_df["总充电功率"] = display_df["显示充电桩数"] * Config.CHARGER_POWER
        chart_title = "扩容演化分析"

    display_areas = display_df["服务区ID"].unique().tolist()
    filtered_stations_df = phase_stations[
        phase_stations["所属服务区"].isin(display_areas)
    ].copy()

    if st.session_state.selected_area is not None:
        if st.session_state.selected_area not in display_areas:
            st.session_state.selected_area = None

    map_center = Config.MAP_CENTER
    if st.session_state.selected_area is not None:
        selected_area_row = display_df[display_df["服务区ID"] == st.session_state.selected_area].iloc[0]
        map_center = [selected_area_row["lat"], selected_area_row["lon"]]

    UIComponents.show_kpi_metrics(filtered_stations_df, display_df, view_mode)

    st.markdown(
        "<div style='height:36px;'></div>",
        unsafe_allow_html=True
    )

    col_map, col_right = st.columns([6, 4])

    with col_map:
        st.markdown(f"### 📍 {chart_title}")

        m = MapGenerator.create_map(
            display_df,
            filtered_stations_df,
            view_mode,
            center=map_center
        )

        map_data = st_folium(
            m,
            width=None,
            height=Config.MAP_HEIGHT,
            returned_objects=["last_object_clicked"],
            key="main_map"
        )

        current_click = map_data.get("last_object_clicked")

        if (
                current_click is not None
                and current_click != st.session_state.last_processed_click
        ):
            st.session_state.last_processed_click = current_click

            need_rerun = MapGenerator.handle_map_click(
                current_click,
                display_df
            )

            if need_rerun:
                st.rerun()

    with col_right:
        col_chart, col_details = st.columns(2)
        with col_chart:
            if st.button(
                    "📊 等级分布",
                    use_container_width=True,
                    type="primary" if st.session_state.display_mode == "chart" else "secondary"
            ):
                st.session_state.display_mode = "chart"
                st.rerun()
        with col_details:
            if st.button(
                    "🎯 服务区详情",
                    use_container_width=True,
                    type="primary" if st.session_state.display_mode == "details" else "secondary"
            ):
                st.session_state.display_mode = "details"
                st.rerun()

        st.markdown("---")

        if st.session_state.display_mode == "chart":
            st.markdown("### 📊 各等级充电站分布")
            UIComponents.create_station_level_chart(filtered_stations_df, view_mode)
        else:
            UIComponents.show_selected_area_details(
                all_service_areas,
                filtered_stations_df
            )


if __name__ == "__main__":
    main()
